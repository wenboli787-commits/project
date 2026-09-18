"""Build an expanded, figure-rich Chinese thesis Results subsection."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_results_section_cn import (
    METRICS,
    add_body_paragraph,
    add_motion_table,
    add_note,
    add_overall_tables,
    configure_page,
    configure_styles,
    metric_means,
    row_for,
    set_cell_margins,
    set_run_font,
    val,
)


OUTPUT_DIR = Path(r"D:\GMR_WORK\GMR\outputs\retargeting_evaluation")
INPUT_CSV = OUTPUT_DIR / "per_motion_metrics.csv"
STATS_CSV = OUTPUT_DIR / "statistical_comparison.csv"
FIGURE_DIR = OUTPUT_DIR / "Figures_Editable_LargeFonts_20260805"
OUTPUT_DOCX = OUTPUT_DIR / "Results_3X_Retargeting_Methods_Academic_CN_v5_FontOnly_EditableFigures.docx"


def pct_reduction(baseline: float, candidate: float) -> float:
    return (baseline - candidate) / baseline * 100.0


def add_figure(doc: Document, image_name: str, caption: str, note: str, width: float = 6.35) -> None:
    image_path = FIGURE_DIR / image_name
    if not image_path.exists():
        raise FileNotFoundError(image_path)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(str(image_path), width=Inches(width))

    cap = doc.add_paragraph(caption, style="Caption")
    cap.paragraph_format.keep_with_next = True
    add_note(doc, note)


def add_lead_summary(doc: Document, df: pd.DataFrame) -> None:
    means = metric_means(df)
    dm = means.loc["direct_mapping"]
    ik = means.loc["basic_ik"]
    gmr = means.loc["gmr"]
    add_body_paragraph(doc, [
        ("主要结果。", True),
        (
            "在 9 个动作、3 种方法的配对评估中，GMR 在 root-relative、尺度归一化、末端执行器和身体姿态误差上均取得最低平均值，"
            f"并将平均 MuJoCo 自碰撞代理率和零容差关节越界率分别降至 {gmr[METRICS['self_collision']] * 100:.2f}% 与 {gmr[METRICS['joint_limit']] * 100:.2f}%。"
            f"Direct Mapping 的 RMS jerk 最低（{dm[METRICS['jerk']]:.2f} rad/s³），Basic IK 的平均脚滑距离最低（{ik[METRICS['foot_sliding']]:.3f} m），"
            "说明三种方法之间存在明确的精度—连续性—接触质量权衡。三种方法在严格质量阈值下均未获得成功样本（0/9），因此本节的结论是相对性能比较，"
            "而非说明任一方法已经生成完全满足物理约束的机器人轨迹。",
            False,
        ),
    ])


def add_scope(doc: Document) -> None:
    doc.add_heading("3.X.1 比较范围与证据组织", level=2)
    add_body_paragraph(doc, [
        ("比较采用统一的机器人模型、30 FPS 评估频率和共同有效时间区间。", True),
        (
            "全部 9 个源动作均分别由 Direct Mapping、Basic IK 和 GMR 生成 Unitree G1 轨迹，形成 27 条方法—动作记录。"
            "为避免只依据平均值判断，本节同时报告逐动作分布、按动作重采样得到的 95% bootstrap 区间、配对非参数检验以及相同评估帧的可视化结果。"
            "其中，误差、脚滑、穿地、MuJoCo 自碰撞代理率、零容差关节越界率和 jerk 均为越低越好。",
            False,
        ),
    ])
    add_body_paragraph(doc, [
        ("动作集合刻意覆盖不同运动学难点。", True),
        (
            "静态站立检验低动态稳定性；行走、跑步、转身和楼梯动作检验根节点轨迹与足部接触；双脚跳和前踢检验快速动态与单腿构型；"
            "05_04 的复杂全身舞蹈动作则用于观察大幅躯干与四肢运动下的关节可达性和全身协调。表 3.X-1 给出动作选择与主要评估目的。",
            False,
        ),
    ])
    add_motion_table(doc)
    motion_table = doc.tables[-1]
    for row in motion_table.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=35, bottom=35, start=100, end=100)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8.2)


def add_motion_preservation(doc: Document, df: pd.DataFrame) -> None:
    means = metric_means(df)
    dm = means.loc["direct_mapping"]
    ik = means.loc["basic_ik"]
    gmr = means.loc["gmr"]

    # Continue naturally after the motion table to avoid a sparse carry-over page.
    doc.add_heading("3.X.2 GMR 的优势主要体现在全身动作保持", level=2)
    add_body_paragraph(doc, [
        ("GMR 在四项运动保持指标上均取得最低的跨动作平均值。", True),
        (
            f"如图 3.X-1 所示，其 root-relative、尺度归一化和末端执行器误差分别为 {gmr[METRICS['root_relative']]:.4f} m、"
            f"{gmr[METRICS['scale_normalised']]:.4f} m 和 {gmr[METRICS['end_effector']]:.4f} m，平均身体姿态误差为 {gmr[METRICS['orientation']]:.2f}°。"
            f"相较 Direct Mapping，上述三项位置误差分别降低 {pct_reduction(dm[METRICS['root_relative']], gmr[METRICS['root_relative']]):.1f}%、"
            f"{pct_reduction(dm[METRICS['scale_normalised']], gmr[METRICS['scale_normalised']]):.1f}% 和 "
            f"{pct_reduction(dm[METRICS['end_effector']], gmr[METRICS['end_effector']]):.1f}%。"
            "逐动作散点也显示，GMR 的分布更集中在低误差区域，而 Direct Mapping 在跑步、楼梯等移动动作上出现了更长的高误差尾部。",
            False,
        ),
    ])
    add_figure(
        doc,
        "Fig_3X_1_Motion_Preservation_Editable.png",
        "图 3.X-1  三种方法在 9 个动作上的运动保持误差分布",
        "注：散点表示单个动作，菱形表示 9 个动作的算术平均值，横线表示对动作进行 20,000 次重采样得到的 95% bootstrap 区间；所有指标越低越好。数据来源：per_motion_metrics.csv。",
    )
    add_body_paragraph(doc, [
        ("配对统计检验支持 GMR 相对 Direct Mapping 的主要优势，但对 Basic IK 的结论需要区分指标。", True),
        (
            "Friedman 检验显示，尺度归一化误差（p = 0.0048）、末端执行器误差（p = 0.0011）和身体姿态误差（p < 0.001）在三种方法间存在显著差异，"
            "而 root-relative 误差未达到显著水平（p = 0.169）。经 Holm 校正后，GMR 相较 Direct Mapping 在尺度归一化、末端执行器和姿态误差上均显著更低；"
            "GMR 相较 Basic IK 则仅在姿态误差上达到显著差异，尺度归一化和末端误差虽具有更低平均值，但应表述为描述性趋势。"
            f"Basic IK 的平均末端误差为 {ik[METRICS['end_effector']]:.4f} m，说明末端约束确实改善了位置匹配，但其平均姿态误差升至 {ik[METRICS['orientation']]:.2f}°，"
            "表明末端位置精度不能替代对整体身体构型的评价。",
            False,
        ),
    ])
    add_overall_tables(doc, df)


def add_physical_results(doc: Document, df: pd.DataFrame) -> None:
    means = metric_means(df)
    dm = means.loc["direct_mapping"]
    ik = means.loc["basic_ik"]
    gmr = means.loc["gmr"]

    # Continue naturally after the summary tables to keep the evidence flow compact.
    doc.add_heading("3.X.3 物理代理指标的改善集中在碰撞与关节可行性", level=2)
    add_body_paragraph(doc, [
        ("GMR 降低了 MuJoCo 自碰撞代理和零容差关节越界，但没有解决所有地面接触问题。", True),
        (
            f"图 3.X-2 显示，GMR 的平均 MuJoCo 自碰撞代理率为 {gmr[METRICS['self_collision']] * 100:.2f}%，低于 Direct Mapping 的 "
            f"{dm[METRICS['self_collision']] * 100:.2f}% 和 Basic IK 的 {ik[METRICS['self_collision']] * 100:.2f}%；其平均零容差关节越界率为 "
            f"{gmr[METRICS['joint_limit']] * 100:.2f}%，而 Direct Mapping 与 Basic IK 分别为 {dm[METRICS['joint_limit']] * 100:.2f}% 和 "
            f"{ik[METRICS['joint_limit']] * 100:.2f}%。尤其是 Basic IK 的零容差关节越界率在 9 个动作上几乎均聚集于 100%，且独立的 0.1° 容差复算仍为 100%，说明逐帧末端匹配经常通过不可行的关节配置实现。GMR 的极小非零零容差结果来自数值精度量级的边界接触，不应解释为实质性超限。",
            False,
        ),
    ])
    add_figure(
        doc,
        "Fig_3X_2_Physical_Plausibility_Editable.png",
        "图 3.X-2  三种方法在 9 个动作上的物理合理性指标",
        "注：最大穿地深度以厘米显示；其他标记含义同图 3.X-1。自碰撞率是 MuJoCo collision geoms 接触代理，不能覆盖全部可见网格穿模；关节越界采用零数值容差。Friedman 检验中，两项代理指标存在显著方法差异，脚滑与最大穿地深度未达到显著水平。数据来源：per_motion_metrics.csv。",
    )
    add_body_paragraph(doc, [
        ("足部接触指标呈现不同于姿态指标的结果。", True),
        (
            f"Basic IK 的平均脚滑距离最低（{ik[METRICS['foot_sliding']]:.4f} m），GMR 为 {gmr[METRICS['foot_sliding']]:.4f} m，"
            f"Direct Mapping 为 {dm[METRICS['foot_sliding']]:.4f} m；但三者差异在 Friedman 检验中不显著（p = 0.264）。"
            f"GMR 的平均最大穿地深度为 {gmr[METRICS['penetration']]:.4f} m，反而高于 Direct Mapping 的 {dm[METRICS['penetration']]:.4f} m 和 "
            f"Basic IK 的 {ik[METRICS['penetration']]:.4f} m，且方法差异同样不显著（p = 0.895）。"
            "脚滑采用脚底高度与垂直速度定义的当前接触掩码，深穿地帧也可能进入该掩码；使用更严格的有效接触或实际 MuJoCo 地面接触时，方法排序会发生变化。"
            "因此，现有证据支持 GMR 在关节可行性代理方面的优势，却不支持依据当前脚滑或穿地口径宣称其地面接触稳定性全面优于两条基线。",
            False,
        ),
    ])
    add_body_paragraph(doc, [
        ("严格成功率进一步强调了这一结论边界。", True),
        (
            "27 条轨迹均完成了评估计算，但每条轨迹至少触发一个预设质量阈值，因此三种方法的严格成功率均为 0/9。"
            "GMR 的失败原因主要集中在持续或较深穿地；Direct Mapping 还经常伴随自碰撞与关节越界；Basic IK 则进一步出现高 jerk 或不连续。"
            "这里的“失败”是质量阈值失败，而不是程序运行失败。",
            False,
        ),
    ])


def add_temporal_results(doc: Document, df: pd.DataFrame) -> None:
    means = metric_means(df)
    dm = means.loc["direct_mapping"]
    ik = means.loc["basic_ik"]
    gmr = means.loc["gmr"]

    doc.add_heading("3.X.4 Basic IK 的末端精度伴随明显时间不连续", level=2)
    add_body_paragraph(doc, [
        ("RMS jerk 揭示了逐帧位置优化未在时间维度上保持稳定。", True),
        (
            f"Direct Mapping、GMR 和 Basic IK 的平均 RMS jerk 分别为 {dm[METRICS['jerk']]:.2f}、{gmr[METRICS['jerk']]:.2f} 和 "
            f"{ik[METRICS['jerk']]:.2f} rad/s³。Basic IK 的数值约为 Direct Mapping 的 {ik[METRICS['jerk']] / dm[METRICS['jerk']]:.2f} 倍，"
            f"也约为 GMR 的 {ik[METRICS['jerk']] / gmr[METRICS['jerk']]:.2f} 倍。图 3.X-3 使用对数横轴保留这一数量级差异，同时仍能展示 Direct Mapping 与 GMR 在不同动作上的离散程度。",
            False,
        ),
    ])
    add_figure(
        doc,
        "Fig_3X_3_Temporal_Consistency_Editable.png",
        "图 3.X-3  三种方法在 9 个动作上的 RMS 关节 jerk 分布",
        "注：横轴采用对数尺度，越低表示关节轨迹越平滑；散点、菱形和区间的含义同图 3.X-1。数据来源：per_motion_metrics.csv。",
    )
    add_body_paragraph(doc, [
        ("统计结果表明三种方法的时间连续性差异并非由单个异常动作驱动。", True),
        (
            "RMS jerk 的 Friedman 检验达到显著水平（p < 0.001）；Holm 校正后的两两比较显示，GMR 显著低于 Basic IK，也显著高于 Direct Mapping（校正后 p = 0.0391）。"
            "这一结果与两种方法的优化方式一致：Direct Mapping 直接继承了较平滑的人体旋转序列，GMR 在全身约束与姿态保持之间进行联合权衡，"
            "而缺少显式时序正则的帧级 Basic IK 更容易在相邻帧切换到不同关节解。该解释属于与观察一致的机制性推断，并不单独构成因果证明。",
            False,
        ),
    ])


def add_same_frame_results(doc: Document, df: pd.DataFrame) -> None:
    doc.add_page_break()
    doc.add_heading("3.X.5 同帧对比揭示了数值差异对应的姿态失真", level=2)
    add_body_paragraph(doc, [
        ("相同评估帧的并列结果使平均指标背后的具体姿态差异可见。", True),
        (
            "图 3.X-4 选择复杂全身动作 05_04 的第 180 帧（6.0 s）和前踢动作 135_04 的第 99 帧（3.3 s）。"
            "六个姿态均从对应 PKL 轨迹重新渲染，并统一采用 Unitree G1 模型、1280×960 分辨率、3.2 m 相机距离、-12° 俯仰角和 180° 世界方位角。"
            "为消除构图差异，渲染前将六帧的根节点位置严格统一为 (0, 0, 0.72 m)，将根节点偏航统一为 0°，并将相机注视点固定为 (0, 0, 0.87 m)；"
            "根节点横滚、俯仰及全部关节坐标保持原轨迹数值不变。因此，同一动作的三列结果具有完全一致的机器人参考位置、相机外参、背景透视和物理成像尺度。"
            "这一刚体归一化仅用于比较局部身体构型；方法之间未经归一化的全局位移、根节点高度和朝向差异仍由原始轨迹上的定量指标评价。",
            False,
        ),
    ])
    add_figure(
        doc,
        "Fig_3X_4_Same_Frame_Comparison_Editable.png",
        "图 3.X-4  两个代表动作在相同评估帧下的三方法对比",
        "注：上行为 05_04 第 180 帧，下行为 135_04 第 99 帧；列顺序为 Direct Mapping、Basic IK 和 GMR。六帧使用同一固定世界相机，并将根节点位置统一为 (0, 0, 0.72 m)、偏航统一为 0°；除该刚体对齐外，根节点横滚、俯仰及关节坐标均未修改。",
    )

    d = row_for(df, "05_04", "direct_mapping")
    b = row_for(df, "05_04", "basic_ik")
    g = row_for(df, "05_04", "gmr")
    add_body_paragraph(doc, [
        ("在 05_04 的大幅全身动作中，三种方法对相同运动阶段给出了明显不同的全身构型。", True),
        (
            "Direct Mapping 保留了侧向抬腿的大致动作幅度，但上、下肢比例与躯干构型受到机器人形态差异影响；Basic IK 为满足末端位置形成了更接近水平展开的过度伸展姿态；"
            "GMR 则在保留大幅腿部运动的同时维持了更连贯的躯干—髋部关系。该观察与定量结果一致：三者的身体姿态误差分别为 "
            f"{val(d, 'orientation'):.2f}°、{val(b, 'orientation'):.2f}° 和 {val(g, 'orientation'):.2f}°，MuJoCo 自碰撞代理率分别为 "
            f"{val(d, 'self_collision') * 100:.2f}%、{val(b, 'self_collision') * 100:.2f}% 和 {val(g, 'self_collision') * 100:.2f}%。"
            f"不过，GMR 的末端执行器误差（{val(g, 'end_effector'):.3f} m）仍高于 Basic IK（{val(b, 'end_effector'):.3f} m），说明较低的整体姿态误差并不等价于最低末端位置误差。",
            False,
        ),
    ])

    d = row_for(df, "135_04", "direct_mapping")
    b = row_for(df, "135_04", "basic_ik")
    g = row_for(df, "135_04", "gmr")
    add_body_paragraph(doc, [
        ("前踢动作中的差异更直接地反映了单腿动态构型的处理能力。", True),
        (
            "在统一根节点位置后，Basic IK 仍表现出明显的双腿屈曲与收缩，削弱了踢腿阶段的动作语义；Direct Mapping 保留了较大的腿部运动，但全身构型更接近跨步或腾空；"
            "GMR 维持了更清晰的支撑腿—摆动腿分工。该动作上，GMR 的 root-relative、尺度归一化和末端执行器误差分别为 "
            f"{val(g, 'root_relative'):.3f} m、{val(g, 'scale_normalised'):.3f} m 和 {val(g, 'end_effector'):.3f} m，均低于 Direct Mapping（"
            f"{val(d, 'root_relative'):.3f} m、{val(d, 'scale_normalised'):.3f} m、{val(d, 'end_effector'):.3f} m）与 Basic IK（"
            f"{val(b, 'root_relative'):.3f} m、{val(b, 'scale_normalised'):.3f} m、{val(b, 'end_effector'):.3f} m）。"
            f"GMR 的姿态误差为 {val(g, 'orientation'):.2f}°，MuJoCo 自碰撞代理率为 {val(g, 'self_collision') * 100:.2f}%，零容差关节越界率为 {val(g, 'joint_limit') * 100:.2f}%，"
            "支持其在该动作上的姿态保持和当前几何代理指标优势，但不能单凭这一帧证明动态平衡能力。",
            False,
        ),
    ])


def add_limitations_and_conclusion(doc: Document) -> None:
    doc.add_heading("3.X.6 结论边界与后续验证", level=2)
    add_body_paragraph(doc, [
        ("综合定量分布、配对检验和同帧结果，GMR 的主要优势是更好地保持全身动作结构并减少机器人不可行姿态。", True),
        (
            "Direct Mapping 具有实现简单和时间轨迹平滑的优点，但难以处理人体与 G1 的身体比例、关节轴和可达范围差异；Basic IK 能够降低部分末端位置误差，"
            "却经常以姿态偏转、关节越界和高 jerk 为代价；GMR 在动作保持、躯干朝向、自碰撞与关节可行性之间取得了更稳定的平衡。"
            "与此同时，脚滑和穿地结果表明，运动学层面的全身重定向仍不能替代接触约束与物理跟踪。",
            False,
        ),
    ])
    add_body_paragraph(doc, [
        ("上述结论仍受五项限制。", True),
        (
            "第一，样本仅包含 9 个动作，且不同动作并非来自随机总体，因此 bootstrap 区间主要用于展示跨动作稳定性，而不应被解释为对所有动作类型的总体置信区间；"
            "第二，评估针对单一 Unitree G1 模型与一组固定参数，尚不能说明结论可直接推广到其他机器人形态；"
            "第三，自碰撞指标仅统计 XML 中启用的 MuJoCo collision geoms，且橡胶手部等可见网格缺少对应碰撞几何，因此不能视为完整的视觉网格穿模率；"
            "第四，脚滑值依赖当前的足部高度与竖直速度接触掩码，零容差关节越界率对数值边界接触敏感；同帧图统一了相机、根节点位置和偏航，仅用于比较局部构型，不能反映原始全局高度、平移或朝向；"
            "第五，本节评价的是运动学重定向轨迹，尚未纳入闭环动力学跟踪误差、执行器限制和真实接触扰动。"
            "后续实验应增加动作数量与重复试次，并在相同物理控制器下比较跟踪成功率、接触冲量、能耗和摔倒率，同时加入足底接触感知优化以检验 GMR 的穿地与脚滑问题能否被系统性降低。",
            False,
        ),
    ])


def build_document() -> Path:
    df = pd.read_csv(INPUT_CSV)
    if len(df) != 27 or df["motion_name"].nunique() != 9 or df["method"].nunique() != 3:
        raise ValueError("Expected 27 rows from 9 motions and 3 methods")
    for column in METRICS.values():
        if not pd.to_numeric(df[column], errors="coerce").notna().all():
            raise ValueError(f"Invalid main metric: {column}")
    if not STATS_CSV.exists():
        raise FileNotFoundError(STATS_CSV)

    doc = Document()
    configure_page(doc)
    configure_styles(doc)
    doc.core_properties.title = "3.X 三种动作重定向方法的学术比较"
    doc.core_properties.subject = "GMR、Direct Mapping 与 Basic IK 的定量及同帧结果比较"
    doc.core_properties.author = ""
    doc.core_properties.keywords = "motion retargeting, GMR, Direct Mapping, Basic IK, Unitree G1"

    doc.add_heading("3.X 三种动作重定向方法的比较", level=1)
    add_lead_summary(doc, df)
    add_scope(doc)
    add_motion_preservation(doc, df)
    add_physical_results(doc, df)
    add_temporal_results(doc, df)
    add_same_frame_results(doc, df)
    add_limitations_and_conclusion(doc)

    add_note(
        doc,
        "数据与复现说明：正文数值来自 per_motion_metrics.csv，显著性结论来自 statistical_comparison.csv；图 3.X-1 至图 3.X-4 由 create_academic_results_figures.py 生成。2026-07-22 在 Ubuntu/WSL 中对 9 个动作、27 条方法轨迹进行了隔离全量复算，主 CSV 与隔离结果逐值一致。",
    )
    doc.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


if __name__ == "__main__":
    print(build_document())
