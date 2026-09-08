from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


OUT = Path("D:/GMR_WORK/thesis_chapter2_method_GMR_MJLab_cn.docx")

CN_FONT = "SimSun"
HEADING_FONT = "SimHei"
LATIN_FONT = "Times New Roman"
ACCENT = RGBColor(31, 77, 120)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(90, 90, 90)


def set_run_font(run, name=CN_FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), LATIN_FONT)
    rfonts.set(qn("w:hAnsi"), LATIN_FONT)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), LATIN_FONT)


def set_style_font(style, east_asia=CN_FONT, latin=LATIN_FONT, size=12, color=None, bold=None):
    style.font.name = latin
    style.font.size = Pt(size)
    if color is not None:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold
    rpr = style._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), latin)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        tbl_grid.append(grid_col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)
    set_run_font(run, size=10, color=MUTED)


def add_para(doc, text, style=None, bold=False, italic=False, color=None, align=None, before=None, after=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, size=12, bold=bold, italic=italic, color=color)
    return p


def add_equation(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, name=LATIN_FONT, size=11.5, color=RGBColor(30, 30, 30))
    return p


def add_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    for run in p.runs:
        set_run_font(run, name=HEADING_FONT, size={1: 16, 2: 13, 3: 12}.get(level, 12), bold=True, color=BLUE if level < 3 else ACCENT)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(item)
        set_run_font(run, size=12)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(item)
        set_run_font(run, size=12)


def add_table(doc, headers, rows, widths_dxa):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths_dxa)
    header_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        set_cell_shading(header_cells[i], "F4F6F9")
        p = header_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(h)
        set_run_font(run, name=HEADING_FONT, size=10.5, bold=True, color=ACCENT)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            if len(text) < 18 and i != 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            set_run_font(run, size=10.5)
    doc.add_paragraph()
    return table


def build_doc():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)

    styles = doc.styles
    set_style_font(styles["Normal"], CN_FONT, LATIN_FONT, 12)
    styles["Normal"].paragraph_format.space_after = Pt(8)
    styles["Normal"].paragraph_format.line_spacing = 1.15
    set_style_font(styles["Heading 1"], HEADING_FONT, LATIN_FONT, 16, BLUE, True)
    styles["Heading 1"].paragraph_format.space_before = Pt(18)
    styles["Heading 1"].paragraph_format.space_after = Pt(10)
    set_style_font(styles["Heading 2"], HEADING_FONT, LATIN_FONT, 13, BLUE, True)
    styles["Heading 2"].paragraph_format.space_before = Pt(12)
    styles["Heading 2"].paragraph_format.space_after = Pt(6)
    set_style_font(styles["Heading 3"], HEADING_FONT, LATIN_FONT, 12, ACCENT, True)
    styles["Heading 3"].paragraph_format.space_before = Pt(8)
    styles["Heading 3"].paragraph_format.space_after = Pt(4)
    set_style_font(styles["List Bullet"], CN_FONT, LATIN_FONT, 12)
    set_style_font(styles["List Number"], CN_FONT, LATIN_FONT, 12)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_run = header.add_run("第二章 Method | GMR 动作重定向与 MJLab 物理仿真跟踪")
    set_run_font(header_run, size=9.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(30)
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run("第二章 方法")
    set_run_font(r, name=HEADING_FONT, size=22, bold=True, color=RGBColor(0, 0, 0))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    r = subtitle.add_run("GMR 动作重定向与 MJLab 物理仿真跟踪方法")
    set_run_font(r, name=HEADING_FONT, size=14, color=MUTED)

    add_para(
        doc,
        "本章描述本文采用的方法体系。研究对象是将人体运动数据转换为 Unitree G1 类人机器人可以表达和执行的参考动作，并在 MJLab/MuJoCo 物理仿真中训练策略对该参考动作进行稳定跟踪。因此，本章重点分为两个部分：第一，General Motion Retargeting（GMR）如何通过人体尺度处理、body 对应关系、位姿目标构造和加权逆运动学生成机器人参考轨迹；第二，MJLab 如何将参考轨迹转化为强化学习跟踪任务，并使用 PD actuator 与 Proximal Policy Optimisation（PPO）训练闭环控制策略。",
    )

    add_heading(doc, "2.1 方法总览与问题定义", 1)
    add_para(
        doc,
        "本文的方法链路从人体动作开始，到机器人在物理仿真中的执行结果结束。GMR 阶段解决的是运动学层面的跨形态映射问题，即人体骨架与 Unitree G1 机器人在身高、肢体比例、自由度数量、关节轴定义和 link 坐标系上均不完全一致，因此不能直接把人体关节角复制给机器人。MJLab 阶段解决的是动力学执行问题，即即使 GMR 产生了几何上合理的机器人姿态序列，具有质量、惯量、接触摩擦和关节力矩限制的机器人仍可能因为脚滑、失衡、速度过大或力矩不足而无法执行该轨迹。",
    )
    add_para(
        doc,
        "整体流程可以概括为：人体动作数据经过 GMR 处理后得到 Unitree G1 的 qpos 参考轨迹；该轨迹被转换为 MJLab 所需的 motion.npz；训练环境在每个控制步读取参考帧，并向策略提供当前机器人状态与参考动作之间的误差信息；PPO 策略输出 29 维关节位置 action；PD actuator 将关节位置目标转换为物理仿真的关节力矩；最后根据 body 位置、姿态、速度、动作平滑性和稳定性计算奖励，并更新策略参数。",
    )
    add_equation(doc, "human motion → GMR retargeting → robot qpos → motion.npz → MJLab tracking → PPO policy")

    add_heading(doc, "2.1.1 状态表示", 2)
    add_para(
        doc,
        "设人体动作序列共有 T 帧。第 t 帧人体动作可以表示为一组人体关键 body 的三维位置和姿态：",
    )
    add_equation(doc, "H_t = {(p_i^h(t), R_i^h(t))}_{i=1}^{N_h},  p_i^h ∈ R^3,  R_i^h ∈ SO(3)")
    add_para(
        doc,
        "其中 p_i^h(t) 表示人体第 i 个 body 在世界坐标系中的位置，R_i^h(t) 表示该 body 的全局旋转矩阵。GMR 的输出是机器人在 MuJoCo 中使用的 generalized coordinate。对 Unitree G1 而言，本文使用 29 DoF 的关节配置，因此每一帧机器人位姿可写为：",
    )
    add_equation(doc, "q_t^r = [p_root(t), q_root(t), q_joint(t)]")
    add_equation(doc, "p_root ∈ R^3,  q_root ∈ H,  q_joint ∈ R^29")
    add_para(
        doc,
        "在实际文件中，GMR/MuJoCo 的 qpos 由 36 个数构成：前三维为 root position，随后四维为 root quaternion，最后 29 维为关节角。该表示保留了机器人在世界坐标系中的全局运动，同时也保留了各个关节的局部配置。",
    )

    add_heading(doc, "2.2 GMR 动作重定向方法", 1)
    add_para(
        doc,
        "GMR 的核心思想是先把人体动作转换为一组机器人 body 的目标位姿，再通过机器人正运动学和逆运动学求解机器人关节配置。它不是简单地对单个关节做一对一映射，而是把 pelvis、torso、腿部、脚部、肩部、肘部和手腕等关键 body 放入统一的优化问题中，通过不同权重同时约束它们的位置和姿态。这样可以在人体与机器人形态不一致时保留动作的整体语义，例如身体朝向、脚端位置、手臂摆动和躯干姿态。",
    )

    add_heading(doc, "2.2.1 人体动作预处理", 2)
    add_para(
        doc,
        "输入人体动作通常来自 AMASS/SMPL-X 风格的数据文件。原始数据包含 pose、translation、shape parameter、mocap framerate 等字段，但这些参数本身不能直接作为机器人 IK 目标。GMR 首先通过人体模型前向运动学恢复逐帧 body 位姿，并提取参与重定向的关键 body。该阶段需要统一帧率、处理人体 root 轨迹、计算全局旋转，并保证每一帧人体 body 的顺序与 GMR 配置文件中的 body 名称一致。",
    )
    add_para(
        doc,
        "人体动作预处理的输出不是人体关节角，而是一组全局 body transform。这样做的好处是，后续 IK 优化只需要关心机器人 link 与人体 body 之间的空间关系，而不依赖人体模型内部的具体关节参数化方式。对于不同来源的人体动作，只要能够得到一致的 body position 和 body orientation，就可以进入同一个 GMR 流程。",
    )

    add_heading(doc, "2.2.2 尺度归一与局部比例调整", 2)
    add_para(
        doc,
        "人体与 Unitree G1 在身高和肢体比例上存在差异。如果直接使用人体关键点的绝对位置作为机器人目标，机器人可能需要达到不可达的脚端或手端位置，导致 IK 失败或关节角落在极限附近。因此，GMR 在 pelvis/root 局部坐标系下对人体 body 位置进行尺度变换。设人体 root 位置为 p_r^h，第 i 个 body 的原始位置为 p_i^h，对应尺度系数为 s_i，则尺度处理可表示为：",
    )
    add_equation(doc, "p̃_r^h = s_r p_r^h")
    add_equation(doc, "p̃_i^h = p̃_r^h + s_i (p_i^h − p_r^h),  i ≠ r")
    add_para(
        doc,
        "该公式说明，除 root 外，各 body 的缩放发生在相对于 root 的局部位移上，而不是对所有全局坐标做同一个比例缩放。这样可以保留人体局部运动结构，同时允许腿部、躯干和手臂使用不同尺度系数。对 Unitree G1 的 SMPL-X 到机器人配置，GMR 使用人体身高假设值，并在实际输入人体高度已知时根据高度比例调整尺度表。该处理使目标动作更接近机器人可达空间，降低 IK 不可行和脚端目标过远的概率。",
    )

    add_heading(doc, "2.2.3 Body 对应关系与目标位姿构造", 2)
    add_para(
        doc,
        "GMR 通过配置文件定义人体 body 与机器人 link 之间的对应关系。以 Unitree G1 为例，pelvis 对应机器人 pelvis，spine3 对应 torso_link，人体左右 hip、knee、foot 对应机器人左右髋、膝、踝或脚端 link，人体左右 shoulder、elbow、wrist 对应机器人左右肩、肘、腕 link。每一组对应关系包含位置权重、姿态权重、位置 offset 和旋转 offset。",
    )
    add_table(
        doc,
        ["人体 body", "机器人 body/link", "作用"],
        [
            ("pelvis", "pelvis", "约束 root 附近的整体平移、身体中心和姿态基准"),
            ("spine3", "torso_link", "约束躯干朝向，使上身动作与人体运动一致"),
            ("left/right foot", "ankle 或 toe 相关 link", "约束脚端位置，影响接触、步态和脚滑"),
            ("left/right shoulder", "shoulder link", "约束上臂根部位置，保持上肢结构"),
            ("left/right elbow", "elbow link", "约束肘部位置，减少手臂姿态失真"),
            ("left/right wrist", "wrist link", "约束手端轨迹，保留挥臂和上肢动作语义"),
        ],
        [1800, 2500, 5060],
    )
    add_para(
        doc,
        "由于人体 body 坐标系与机器人 link 坐标系定义不同，即使两者表示的是相似的身体部位，其局部坐标轴也不一定一致。例如人体脚部坐标系、机器人 ankle link 或 toe link 坐标系的前向轴和竖直轴可能不同。因此，GMR 对每个目标 body 使用旋转 offset 和位置 offset。设配置中第 i 个 body 的位置 offset 为 o_i，旋转 offset 为 R_i^off，则目标姿态和目标位置可写为：",
    )
    add_equation(doc, "R_i* = R_i^h R_i^off")
    add_equation(doc, "p_i* = p̃_i^h + R_i* o_i")
    add_para(
        doc,
        "其中 R_i* 和 p_i* 是最终传递给 IK task 的目标位姿。旋转 offset 使人体和机器人对应 body 的坐标轴语义对齐；位置 offset 则补偿 body 原点位置的差异。例如人体 wrist 的原点可能位于手腕关节附近，而机器人 wrist link 的几何原点可能偏向 link 中心，因此需要一个局部偏移量来使目标点更合理。",
    )

    add_heading(doc, "2.2.4 加权逆运动学优化", 2)
    add_para(
        doc,
        "当所有目标 body 位姿构造完成后，GMR 通过逆运动学求解机器人配置 q。设机器人第 k 个 body 的正运动学位置和姿态分别为 p_k^r(q) 和 R_k^r(q)，目标位置和姿态分别为 p_k* 和 R_k*。GMR 的目标是在关节限制内寻找 q，使所有被跟踪 body 的位置误差和姿态误差加权最小。",
    )
    add_equation(
        doc,
        "min_{Δq} Σ_{k∈B} w_k^p ||p_k^r(q+Δq) − p_k*||_2^2 + w_k^R ||Log((R_k*)^T R_k^r(q+Δq))||_2^2 + λ||Δq||_2^2",
    )
    add_para(
        doc,
        "其中 B 是参与匹配的机器人 body 集合，w_k^p 和 w_k^R 分别是位置与姿态权重，Log(·) 是 SO(3) 到李代数 so(3) 的映射，用于把旋转矩阵误差转换为三维轴角误差向量。阻尼项 λ||Δq||² 用于抑制过大的关节更新，提高奇异姿态和目标不可达情况下的数值稳定性。",
    )
    add_para(
        doc,
        "在每一次迭代中，IK 求解器会根据当前机器人配置计算各个 body 的雅可比矩阵。将误差向量记为 e(q)，雅可比记为 J(q)，局部线性化后有：",
    )
    add_equation(doc, "e(q+Δq) ≈ e(q) + J(q)Δq")
    add_para(
        doc,
        "若暂时忽略不等式约束，阻尼加权最小二乘问题的更新方向可写为：",
    )
    add_equation(doc, "Δq = −(J^T W J + λI)^{-1} J^T W e")
    add_para(
        doc,
        "实际实现中还需要考虑关节角限制，因此 GMR 使用带约束的 IK/QP 求解器。关节限制可写为：",
    )
    add_equation(doc, "q_min ≤ q + Δq ≤ q_max")
    add_para(
        doc,
        "如果启用速度限制，还可以加入相邻帧之间的速度约束：",
    )
    add_equation(doc, "|(q_t − q_{t−1}) / Δt| ≤ q_dot,max")
    add_para(
        doc,
        "本项目使用的 GMR 实现基于 mink.solve_ik，求解器为 daqp，阻尼系数为 0.5，每一帧最多迭代 10 次。当相邻两次迭代的误差下降量小于 0.001 时，求解提前停止。这种设置在计算效率和 IK 精度之间取得平衡：每一帧不需要求到完全零误差，但需要得到一个满足机器人关节限制且能保持连续性的可用姿态。",
    )

    add_heading(doc, "2.2.5 两阶段 GMR 求解策略", 2)
    add_para(
        doc,
        "GMR 并不是把所有 body task 一次性放入同一个优化问题中求解，而是使用两阶段匹配表。第一阶段通常强调 root、脚端和躯干等决定整体姿态与接触结构的 body，使机器人先获得稳定的全身配置；第二阶段再加入 hip、knee、shoulder、elbow、wrist 等局部 body 约束，进一步细化肢体姿态。两阶段流程可表示为：",
    )
    add_equation(doc, "q_t^(1) = IK_1(q_{t−1}, T_1(t))")
    add_equation(doc, "q_t^(2) = IK_2(q_t^(1), T_2(t))")
    add_equation(doc, "q_t^r = q_t^(2)")
    add_para(
        doc,
        "其中 T_1(t) 和 T_2(t) 分别表示第 t 帧第一阶段和第二阶段的目标集合。由于第 t 帧通常以上一帧结果 q_{t−1} 作为初值，GMR 具有天然的时间连续性。若每帧都从默认站立姿态重新求解，IK 可能会在多个局部解之间跳变；而逐帧继承初值可以减少不必要的关节突变，使输出轨迹更适合后续物理跟踪。",
    )

    add_heading(doc, "2.2.6 GMR 输出轨迹与可行性含义", 2)
    add_para(
        doc,
        "GMR 的输出是机器人运动学参考，而不是已经经过动力学验证的可执行控制轨迹。对 Unitree G1，每一帧 qpos 的结构为：",
    )
    add_equation(doc, "qpos_t = [x, y, z, q_w, q_x, q_y, q_z, q_1, q_2, ..., q_29]")
    add_para(
        doc,
        "其中 root quaternion 使用 MuJoCo 常见的 wxyz 顺序，关节角顺序必须与 Unitree G1 MJCF 模型一致。标准输出文件还会保存 qvel、root_pos、root_quat、dof_pos、body_pos、body_quat、contact 和 fall 等字段。GMR 的成功说明该动作在运动学上可以被机器人模型近似表达，但不保证机器人在物理仿真中能够稳定完成。例如，一个几何姿态序列可能脚端位置正确，但质心移动过快、支撑相不足或关节速度过高，导致 MJLab 中的机器人摔倒。因此，GMR 输出需要进入 MJLab 做进一步动力学跟踪验证。",
    )

    add_heading(doc, "2.3 从 GMR 轨迹到 MJLab 训练数据", 1)
    add_para(
        doc,
        "MJLab 的 tracking 任务不直接读取 GMR 的 pkl 文件，而是使用 motion.npz 作为参考动作数据库。因此，GMR 输出需要经过格式转换。转换过程首先从 qpos 中提取 root position、root quaternion 和 29 维关节角，然后写入 MJLab CSV。CSV 每一帧包含 36 列：前三列为 root position，随后四列为 root quaternion，最后 29 列为 joint position。",
    )
    add_equation(doc, "root_pos = qpos[:, 0:3]")
    add_equation(doc, "root_quat_wxyz = qpos[:, 3:7]")
    add_equation(doc, "dof_pos = qpos[:, 7:36]")
    add_para(
        doc,
        "需要特别注意四元数顺序。GMR/MuJoCo qpos 通常使用 wxyz，而 MJLab CSV 转换阶段使用 xyzw。因此转换时需要执行：",
    )
    add_equation(doc, "[q_x, q_y, q_z, q_w] = [qpos_4, qpos_5, qpos_6, qpos_3]")
    add_para(
        doc,
        "CSV 之后会通过 MJLab 的转换工具生成 motion.npz。该工具在 MuJoCo 中播放参考关节轨迹，并通过机器人前向运动学计算每个 body 的世界坐标位置、姿态、线速度和角速度。motion.npz 中的主要字段包括 fps、joint_pos、joint_vel、body_pos_w、body_quat_w、body_lin_vel_w 和 body_ang_vel_w。后续 reward 与 observation 大量依赖 body 级别信息，因此 body ordering、关节顺序和 quaternion convention 必须完全正确，否则策略会学习到错误的跟踪目标。",
    )
    add_para(
        doc,
        "本文的 GMR 轨迹通常以 30 FPS 保存，而 MJLab tracking 环境的 policy step 为 0.02 s，即 50 Hz。因此转换阶段需要进行时间重采样，使 reference motion 与策略控制频率一致。若不重采样，策略在训练时读取的参考速度和相邻帧位移会与仿真步长不匹配，从而影响速度奖励、姿态变化节奏和动作稳定性。",
    )

    add_heading(doc, "2.4 MJLab 物理仿真环境原理", 1)
    add_para(
        doc,
        "MJLab 是基于 MuJoCo 的机器人学习框架。本文使用的任务为 Mjlab-Tracking-Flat-Unitree-G1，机器人模型为 Unitree G1，地形为平地 plane，控制方式为 JointPositionAction，训练算法为 RSL-RL PPO。MuJoCo 负责计算机器人刚体动力学、关节约束、接触力、摩擦和 actuator 输出，MJLab 则在其上封装 observation、action、reward、termination 和 motion command。",
    )
    add_para(
        doc,
        "在广义坐标下，机器人动力学可概括为：",
    )
    add_equation(doc, "M(q)q¨ + C(q,q˙)q˙ + g(q) = S^T τ + J_c(q)^T f_c")
    add_para(
        doc,
        "其中 M(q) 是质量矩阵，C(q,q˙)q˙ 表示科氏力和离心项，g(q) 是重力项，τ 是关节 actuator 产生的力矩，J_c^T f_c 是接触约束力对广义坐标的作用。MuJoCo 在每个物理步中求解该动力学系统，并处理脚底与地面的非穿透约束和摩擦。与纯运动学 GMR 不同，MJLab 中的机器人必须在这些动力学条件下保持平衡和接触稳定。",
    )
    add_para(
        doc,
        "本任务的仿真 timestep 为 0.005 s，decimation 为 4。因此 MuJoCo 每 0.005 s 积分一次，而策略每 4 个物理步输出一次 action：",
    )
    add_equation(doc, "Δt_policy = 0.005 × 4 = 0.02 s")
    add_para(doc, "这对应 50 Hz 的策略控制频率。环境 episode 默认长度为 10 s。")

    add_heading(doc, "2.4.1 跟踪 body 与 anchor", 2)
    add_para(
        doc,
        "MJLab tracking 不只跟踪 root，而是同时跟踪多个关键 body。Unitree G1 配置中，torso_link 被设置为 anchor body。被跟踪的 body 包括 pelvis、左右 hip roll link、左右 knee link、左右 ankle roll link、torso_link、左右 shoulder roll link、左右 elbow link 和左右 wrist yaw link。这些 body 覆盖了骨盆、腿部、躯干和上肢，能够评价 whole-body tracking，而不是只评价根节点轨迹。",
    )
    add_para(
        doc,
        "Anchor 的作用是建立参考动作与当前机器人之间的相对坐标关系。若直接在世界坐标中比较所有 body，机器人整体平移或 heading 上的小偏差会被放大，导致 reward 过度惩罚已经保持局部姿态的状态。因此，MJLab 会把参考 body 位置转换到以当前机器人 anchor 为基准的局部表示，使 reward 和 observation 更关注相对于躯干的全身结构误差。",
    )
    add_equation(doc, "p_a,b^ref = (R_a^robot)^T (p_a^ref − p_a^robot)")
    add_equation(doc, "R_a,b^ref = (R_a^robot)^T R_a^ref")

    add_heading(doc, "2.4.2 MotionCommand 与参考帧采样", 2)
    add_para(
        doc,
        "MotionCommand 是 MJLab 中管理参考动作的模块。它为每个并行环境维护一个当前参考帧索引 τ_e。每个 policy step，环境根据 τ_e 从 motion.npz 中读取参考 joint_pos、joint_vel、body_pos_w、body_quat_w、body_lin_vel_w 和 body_ang_vel_w，并在 step 结束后推进到下一帧。",
    )
    add_equation(doc, "τ_e ∈ {0, 1, ..., T−1}")
    add_para(
        doc,
        "训练时通常不总是从动作开头 reset。MJLab 使用起始帧采样，使不同并行环境从动作的不同片段开始训练，从而提升样本多样性。对于容易失败的片段，adaptive sampling 会提高其被采样的概率。设第 b 个时间 bin 的历史失败统计为 C_b，当前 rollout 中的失败统计为 Ĉ_b，则可用指数滑动形式表示为：",
    )
    add_equation(doc, "C_b ← α Ĉ_b + (1−α)C_b")
    add_equation(doc, "P(b) = (C_b + ε) / Σ_j(C_j + ε)")
    add_para(
        doc,
        "其中 ε 保留均匀采样成分，避免训练完全忽略当前看起来较容易的片段。该机制对于复杂动作尤其重要，因为策略如果只在简单片段上获得高 reward，可能无法学会从高速度、单脚支撑或上身大幅摆动等困难状态中恢复。",
    )

    add_heading(doc, "2.4.3 Reset 随机化与鲁棒性", 2)
    add_para(
        doc,
        "在每次 reset 时，环境会根据参考帧初始化机器人状态，并加入一定随机扰动。扰动可以作用于 root 位置、root 姿态、root 速度和关节位置。其目的不是改变参考动作本身，而是让策略学习从偏离参考的状态回到参考轨迹附近。简化地说，初始状态可写为：",
    )
    add_equation(doc, "x_0 = x_0^ref + ε_x")
    add_para(
        doc,
        "此外，MJLab 还可以加入 domain randomization，例如 base center of mass 偏移、encoder bias、脚底摩擦变化以及随机外力或速度扰动。这些随机化提高策略对模型误差、接触变化和传感器偏差的鲁棒性，使策略不只适应一个理想的确定性仿真环境。",
    )

    add_heading(doc, "2.5 Observation、Action 与 PD Actuator", 1)
    add_heading(doc, "2.5.1 Actor observation", 2)
    add_para(
        doc,
        "PPO actor 接收的 observation 同时包含参考动作信息和机器人当前状态。其核心思想是让策略知道“当前机器人在哪里”“参考动作期望机器人在哪里”以及“上一时刻策略输出了什么”。Actor observation 主要包括当前参考 joint_pos 与 joint_vel、参考 anchor 相对机器人 anchor 的位置和姿态、base linear velocity、base angular velocity、当前关节位置和速度、以及上一时刻 action。",
    )
    add_equation(doc, "o_t^π = [q_ref, q˙_ref, p_anchor^b, R_anchor^b, v_base, ω_base, q_rel, q˙_rel, a_{t−1}] + η")
    add_para(
        doc,
        "其中 η 表示训练时加入的 observation noise。对 anchor orientation，MJLab 使用相对旋转矩阵的前两列作为 observation，而不是直接使用四元数。这可以避免四元数 q 与 −q 表示同一姿态导致的符号不连续问题。Actor observation 被设计得相对接近真实机器人可获得的信息，使训练得到的策略更接近可部署的状态反馈控制器。",
    )

    add_heading(doc, "2.5.2 Critic observation", 2)
    add_para(
        doc,
        "Critic 在训练阶段可以使用比 actor 更完整的状态信息。除 actor observation 外，critic 还可访问机器人 body position、body orientation 等 privileged information。这样做不会改变策略部署时的输入，但可以提高 value function 对状态好坏的估计精度，从而降低 advantage estimation 的方差并提升 PPO 训练稳定性。",
    )

    add_heading(doc, "2.5.3 Action 到关节目标", 2)
    add_para(
        doc,
        "本文任务使用 JointPositionAction。策略输出不是关节力矩，而是 29 维归一化 action。MJLab 将 raw action 经过 scale 和 offset 转换为关节位置目标：",
    )
    add_equation(doc, "a_t ∈ R^29")
    add_equation(doc, "q_t^des = q^default + S ⊙ a_t")
    add_para(
        doc,
        "其中 S 是每个关节对应的 action scale，q^default 是机器人默认站立姿态下的关节位置。若训练中加入 encoder bias，实际写入 actuator 的目标为：",
    )
    add_equation(doc, "q_t^target = q_t^des − b_enc")
    add_para(
        doc,
        "这种 action 设计降低了强化学习难度。策略不需要直接学习高频力矩，而是学习下一步每个关节应趋向的目标角度；底层 PD actuator 再根据关节误差生成力矩。",
    )

    add_heading(doc, "2.5.4 PD actuator 原理", 2)
    add_para(
        doc,
        "MJLab 中的 position action 最终通过 PD actuator 转换为关节力矩。对第 j 个关节，理想 PD 控制律为：",
    )
    add_equation(doc, "τ_j = K_p,j(q_j^target − q_j) + K_d,j(q˙_j^target − q˙_j) + τ_j^ff")
    add_para(
        doc,
        "对所有关节写成向量形式：",
    )
    add_equation(doc, "τ = K_p(q^target − q) + K_d(q˙^target − q˙) + τ^ff")
    add_para(
        doc,
        "计算得到的力矩会被 actuator limit 限制：",
    )
    add_equation(doc, "τ^clip = clip(τ, −τ_max, τ_max)")
    add_para(
        doc,
        "PD actuator 的刚度 K_p 决定关节朝目标角度收敛的强度，阻尼 K_d 抑制速度误差和振荡。若 K_p 过低，机器人无法快速跟随参考动作；若 K_p 过高，系统可能出现抖动或接触冲击。PPO 策略学习的是 q_target 的动态调节，而不是替代 PD 控制器本身。",
    )

    add_heading(doc, "2.6 MJLab Reward 设计", 1)
    add_para(
        doc,
        "MJLab tracking reward 的目标是使机器人在物理仿真中尽量接近 GMR 参考动作，同时保持平滑、稳定且不违反关节和碰撞约束。总奖励可以写为：",
    )
    add_equation(doc, "r_t = Σ_i w_i r_i(t) + Σ_j λ_j c_j(t)")
    add_para(
        doc,
        "其中 r_i 是正向跟踪奖励，c_j 是惩罚项。主要正向奖励采用指数型形式：误差越小，奖励越接近 1；误差增大时，奖励快速下降。指数奖励常用于 motion tracking，因为它既强调精确跟踪，又不会像线性误差那样在大误差时无限增大惩罚幅度。",
    )

    add_heading(doc, "2.6.1 Root/anchor 跟踪奖励", 2)
    add_para(
        doc,
        "Root 或 anchor 位置奖励衡量参考 anchor 与机器人 anchor 的位置差：",
    )
    add_equation(doc, "e_root,pos = ||p_a^ref − p_a^robot||_2^2")
    add_equation(doc, "r_root,pos = exp(−e_root,pos / σ_p^2)")
    add_para(
        doc,
        "姿态奖励使用四元数或旋转距离表示：",
    )
    add_equation(doc, "r_root,ori = exp(−d_q(q_a^ref, q_a^robot)^2 / σ_R^2)")
    add_para(
        doc,
        "在本文配置中，root position 的 σ_p 为 0.3、权重为 0.5；root orientation 的 σ_R 为 0.4、权重为 0.5。该部分奖励保证机器人整体身体中心和躯干朝向不会偏离参考轨迹太远。",
    )

    add_heading(doc, "2.6.2 Body 位置、姿态和速度奖励", 2)
    add_para(
        doc,
        "只跟踪 root 会导致四肢动作失真，因此 MJLab 还对多个关键 body 的位置、姿态和速度进行跟踪。设跟踪 body 数量为 N_b，则 body position error 可写为：",
    )
    add_equation(doc, "e_body,pos = (1/N_b) Σ_{k=1}^{N_b} ||p_k^ref,rel − p_k^robot||_2^2")
    add_equation(doc, "r_body,pos = exp(−e_body,pos / σ_bp^2)")
    add_para(
        doc,
        "Body orientation reward 对多个 body 的旋转误差求平均：",
    )
    add_equation(doc, "e_body,ori = (1/N_b) Σ_{k=1}^{N_b} d_q(q_k^ref,rel, q_k^robot)^2")
    add_equation(doc, "r_body,ori = exp(−e_body,ori / σ_bR^2)")
    add_para(
        doc,
        "速度奖励用于约束动作节奏，而不仅仅是静态姿态：",
    )
    add_equation(doc, "r_lin = exp(−[(1/N_b)Σ_k ||v_k^ref − v_k^robot||_2^2] / σ_v^2)")
    add_equation(doc, "r_ang = exp(−[(1/N_b)Σ_k ||ω_k^ref − ω_k^robot||_2^2] / σ_ω^2)")
    add_para(
        doc,
        "本文配置中，body position 的 σ 为 0.3、body orientation 的 σ 为 0.4、body linear velocity 的 σ 为 1.0、body angular velocity 的 σ 为 3.14，四项权重均为 1.0。这些项共同使策略同时关注姿态形状、全身运动节奏和局部肢体速度。",
    )

    add_heading(doc, "2.6.3 平滑性、关节限制与自碰撞惩罚", 2)
    add_para(
        doc,
        "动作变化率惩罚用于抑制策略输出抖动：",
    )
    add_equation(doc, "c_action = ||a_t − a_{t−1}||_2^2")
    add_para(
        doc,
        "该项作用于 raw action，因此直接鼓励神经网络输出连续的关节目标。关节限位惩罚计算关节超过 soft joint limit 的程度：",
    )
    add_equation(doc, "c_limit = Σ_j max(q_j,min − q_j, 0) + max(q_j − q_j,max, 0)")
    add_para(
        doc,
        "自碰撞惩罚通过 contact sensor 检测机器人内部不合理碰撞。当某个接触力超过阈值 f_th 时计入惩罚：",
    )
    add_equation(doc, "c_collision = Σ_h 1(||f_h|| > f_th)")
    add_para(
        doc,
        "在本文配置中，action rate penalty 权重为 -0.1，joint limit penalty 权重为 -10.0，自碰撞惩罚权重为 -10.0，碰撞力阈值为 10.0。这些惩罚项限制策略通过不自然关节角、剧烈动作或肢体穿插来获得短期跟踪奖励。",
    )

    add_table(
        doc,
        ["奖励或惩罚项", "数学含义", "配置含义"],
        [
            ("root position", "anchor 位置误差的指数奖励", "σ=0.3，weight=0.5"),
            ("root orientation", "anchor 姿态误差的指数奖励", "σ=0.4，weight=0.5"),
            ("body position", "关键 body 相对位置误差", "σ=0.3，weight=1.0"),
            ("body orientation", "关键 body 姿态误差", "σ=0.4，weight=1.0"),
            ("body linear velocity", "关键 body 线速度误差", "σ=1.0，weight=1.0"),
            ("body angular velocity", "关键 body 角速度误差", "σ=3.14，weight=1.0"),
            ("action rate", "相邻 action 差值平方", "weight=-0.1"),
            ("joint limit", "超过 soft joint limit 的距离", "weight=-10.0"),
            ("self collision", "超过阈值的内部接触", "threshold=10.0，weight=-10.0"),
        ],
        [2000, 4200, 3160],
    )

    add_heading(doc, "2.7 Episode 终止条件", 1)
    add_para(
        doc,
        "除正常 time-out 外，MJLab tracking 环境会在机器人明显偏离参考或失去稳定性时提前终止 episode。提前终止不是单纯为了节省计算，而是向 PPO 提供清晰的失败信号，使策略避免进入不可恢复状态。主要终止条件包括 anchor 高度偏差过大、anchor 姿态偏差过大，以及末端 body 高度偏差过大。",
    )
    add_equation(doc, "|z_a^ref − z_a^robot| > 0.25")
    add_equation(doc, "∃ k∈E,  |z_k^ref,rel − z_k^robot| > 0.25")
    add_para(
        doc,
        "其中 E 包括左右脚踝和左右手腕等末端 body。姿态终止通过比较参考 anchor 和机器人 anchor 下的 projected gravity 完成，当姿态偏差超过阈值 0.8 时认为机器人已经明显失衡。对于 motion tracking，终止条件能够反映 reference 是否可执行：如果某些片段频繁触发终止，通常说明该片段存在动力学难点，例如支撑相不足、root 高度变化过快、脚端目标不合理或上身动作导致质心偏移过大。",
    )

    add_heading(doc, "2.8 PPO 在 MJLab 中的训练原理", 1)
    add_para(
        doc,
        "本文使用 PPO 训练 Unitree G1 的 tracking policy。PPO 是 on-policy actor-critic 方法，适合连续控制任务。它通过在当前策略下采样 rollout，估计每个状态动作对的 advantage，再使用 clipped surrogate objective 更新策略。与直接最大化策略梯度不同，PPO 限制新旧策略之间的变化幅度，从而提高训练稳定性。",
    )

    add_heading(doc, "2.8.1 MDP 表示", 2)
    add_para(
        doc,
        "MJLab tracking 任务可以表示为马尔可夫决策过程：",
    )
    add_equation(doc, "(S, A, P, r, γ)")
    add_para(
        doc,
        "其中状态 S 对应 actor/critic observation，动作 A 为 29 维关节位置 action，状态转移 P 由 MuJoCo 动力学、接触求解、PD actuator 和随机化共同决定，奖励 r 由 tracking reward 和 penalty 构成，γ 是折扣因子。策略为高斯策略：",
    )
    add_equation(doc, "a_t ~ π_θ(a_t | s_t)")
    add_para(
        doc,
        "Actor 网络输出动作分布，critic 网络估计状态价值：",
    )
    add_equation(doc, "V_φ(s_t) ≈ E[Σ_{k=0}^{∞} γ^k r_{t+k}]")

    add_heading(doc, "2.8.2 Advantage estimation", 2)
    add_para(
        doc,
        "PPO 需要估计当前 action 相对于平均策略表现的优势。首先计算 TD residual：",
    )
    add_equation(doc, "δ_t = r_t + γV_φ(s_{t+1}) − V_φ(s_t)")
    add_para(
        doc,
        "随后使用 generalized advantage estimation（GAE）：",
    )
    add_equation(doc, "Â_t = Σ_{l=0}^{∞} (γλ)^l δ_{t+l}")
    add_para(
        doc,
        "本文配置中 γ=0.99，λ=0.95。γ 越大，策略越重视长期 reward；λ 控制 bias-variance trade-off，较高 λ 可以利用更长时间范围的回报信息，但方差也更高。",
    )

    add_heading(doc, "2.8.3 PPO clipped objective", 2)
    add_para(
        doc,
        "设旧策略为 π_θold，新策略为 π_θ，概率比为：",
    )
    add_equation(doc, "ρ_t(θ) = π_θ(a_t|s_t) / π_θold(a_t|s_t)")
    add_para(
        doc,
        "PPO 的 clipped surrogate objective 为：",
    )
    add_equation(
        doc,
        "L^CLIP(θ) = E_t[min(ρ_t(θ)Â_t, clip(ρ_t(θ), 1−ε, 1+ε)Â_t)]",
    )
    add_para(
        doc,
        "本文使用 ε=0.2。当新策略相对旧策略变化过大时，clip 项会阻止目标函数继续鼓励该方向的更新，从而避免策略因为一次过大的梯度更新而崩溃。对于 humanoid tracking，策略崩溃通常表现为突然输出极端关节目标、机器人摔倒、reward 大幅下降，因此 clipping 对训练稳定性非常重要。",
    )

    add_heading(doc, "2.8.4 Value loss 与 entropy bonus", 2)
    add_para(
        doc,
        "Critic 通过 value loss 学习回报估计：",
    )
    add_equation(doc, "L^V(φ) = E_t[(V_φ(s_t) − R̂_t)^2]")
    add_para(
        doc,
        "策略还包含 entropy bonus，用于鼓励探索，避免过早收敛到动作方差过小的局部策略。综合损失可以写为：",
    )
    add_equation(doc, "L(θ,φ) = −L^CLIP(θ) + c_v L^V(φ) − c_H H(π_θ)")
    add_para(
        doc,
        "本文配置中 value_loss_coef 为 1.0，entropy_coef 为 0.005，desired_kl 为 0.01，max_grad_norm 为 1.0。desired KL 用于 adaptive learning rate schedule：当新旧策略 KL divergence 过大时降低学习率，当更新过小时提高或保持学习率，从而让策略更新维持在合适范围内。",
    )

    add_heading(doc, "2.8.5 网络结构与训练批次", 2)
    add_para(
        doc,
        "Unitree G1 tracking 的 actor 和 critic 均使用多层感知机。隐藏层维度为 512、256、128，激活函数为 ELU，observation normalization 开启，actor 初始标准差为 1.0，学习率为 1e-3 并使用 adaptive schedule。每次 iteration，每个并行环境采样 24 个 policy step，然后对收集到的 rollout 做 5 个 learning epochs，并分成 4 个 mini-batches 更新。",
    )
    add_table(
        doc,
        ["参数", "数值", "说明"],
        [
            ("hidden dims", "512, 256, 128", "actor 与 critic 的 MLP 隐藏层结构"),
            ("activation", "ELU", "非线性激活函数"),
            ("init std", "1.0", "高斯 actor 初始动作标准差"),
            ("learning rate", "1e-3", "adaptive schedule 下的初始学习率"),
            ("gamma", "0.99", "折扣因子"),
            ("lambda", "0.95", "GAE 参数"),
            ("clip param", "0.2", "PPO clipping 范围"),
            ("entropy coef", "0.005", "探索奖励权重"),
            ("desired KL", "0.01", "自适应学习率目标 KL"),
            ("num steps per env", "24", "每次 rollout 每个环境采样步数"),
            ("learning epochs", "5", "每轮数据重复优化次数"),
            ("mini-batches", "4", "每轮数据划分的 mini-batch 数量"),
            ("max iterations", "30000", "最大训练迭代次数"),
        ],
        [2300, 1900, 5160],
    )
    add_para(
        doc,
        "若使用 512 个并行环境，则每次 iteration 收集 512×24=12288 个 transition；若使用 4096 个并行环境，则每次 iteration 收集 4096×24=98304 个 transition。并行环境越多，采样效率越高，但 GPU/内存占用也越大。实际实验中可以先使用较少环境进行 smoke test，确认 motion.npz、reward、termination 和 observation 正常后，再扩大并行环境数量进行正式训练。",
    )

    add_heading(doc, "2.9 评价指标", 1)
    add_para(
        doc,
        "本文的评价需要同时覆盖 GMR 参考轨迹质量和 MJLab 物理跟踪质量。GMR 评价关注机器人参考姿态是否合理、是否满足关节限制、是否存在脚滑、穿地或突然跳变；MJLab 评价关注机器人在动力学环境中是否能稳定跟踪该参考轨迹。",
    )
    add_heading(doc, "2.9.1 GMR 参考轨迹评价", 2)
    add_para(
        doc,
        "GMR 参考轨迹的常用指标包括 key-body 位置误差、root-relative body error、bone orientation error、foot sliding、ground penetration、self collision、sudden jump 和 success rate。Mean per-keybody position error 可写为：",
    )
    add_equation(doc, "MPKPE = (1/N)Σ_{i=1}^{N} ||p_i^ref − p_i^robot||_2")
    add_para(
        doc,
        "Root-relative MPKPE 去除全局平移和 heading 误差，更关注机器人身体局部结构是否与参考一致：",
    )
    add_equation(doc, "R-MPKPE = (1/N)Σ_{i=1}^{N} ||p_i^ref,rel − p_i^robot||_2")
    add_para(
        doc,
        "脚滑指标衡量脚在接触地面时的水平位移。若脚部高度低于接触阈值 h_c，则接触阶段脚端水平移动可累计为：",
    )
    add_equation(doc, "D_slip = Σ_t 1(z_foot(t)<h_c) ||p_foot^xy(t+1) − p_foot^xy(t)||_2")
    add_para(
        doc,
        "这些指标能够定位 GMR 参考轨迹中的具体问题。例如 MPKPE 较低但 foot sliding 较高，说明 body 姿态大体正确但脚部接触不稳定；R-MPKPE 较低但 raw global error 较高，说明局部姿态保持良好但全局 root 轨迹或 heading 存在偏移。",
    )

    add_heading(doc, "2.9.2 MJLab 跟踪评价", 2)
    add_para(
        doc,
        "MJLab 跟踪评价关注物理执行结果。主要指标包括 episode return、episode length、root tracking error、body MPKPE、R-MPKPE、end-effector error、joint velocity error、fall rate、success rate、action smoothness、joint limit penalty 和 self-collision penalty。若策略能够在较长 episode 中保持低 body error、低末端误差和低 fall rate，则说明 GMR 参考轨迹不仅在运动学上合理，也能被物理控制器较稳定地执行。",
    )
    add_para(
        doc,
        "需要注意的是，GMR 误差和 MJLab 跟踪误差代表不同层面的失败。GMR 误差高通常说明人体到机器人形态映射本身不够好；MJLab 跟踪失败则可能来自动力学难度、PD 参数、reward 设计、策略训练不足或参考轨迹本身不可执行。论文实验分析时应把这两类误差分开讨论，避免把控制失败简单归因于重定向失败，或把重定向问题误认为 PPO 没有收敛。",
    )

    add_heading(doc, "2.10 本章小结", 1)
    add_para(
        doc,
        "本章详细说明了本文的核心方法。GMR 部分通过人体动作预处理、尺度归一、body 对应关系、offset 目标构造和加权逆运动学，把人体动作转换为 Unitree G1 的 qpos 参考轨迹。MJLab 部分将该轨迹转换为 motion.npz，在 MuJoCo 中建立物理跟踪任务，并通过 observation、JointPositionAction、PD actuator、tracking reward、termination 和 adaptive sampling 构成强化学习训练环境。PPO 通过 actor-critic、GAE 和 clipped objective 学习闭环策略，使机器人在动力学约束下尽量跟随 GMR 参考动作。后续实验章节可以基于本章定义的指标，分别分析参考轨迹质量、物理执行能力和失败案例。",
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    output = build_doc()
    print(output)
